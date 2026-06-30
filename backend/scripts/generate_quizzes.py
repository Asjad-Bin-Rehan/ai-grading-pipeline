"""Generate varied student quiz DOCX files by applying simple text substitutions to existing templates.

Usage:
    python generate_quizzes.py --count 10 --outdir ../../samples/generated

This script requires `python-docx` to be installed in the project's venv.
"""
import random
import re
from pathlib import Path
from docx import Document
import argparse

ROOT = Path(__file__).resolve().parents[2]
SAMPLES_DIR = ROOT / 'samples'
OUT_DIR = SAMPLES_DIR / 'generated'

# Simple substitution rules (pattern -> list of replacements)
SUBSTITUTIONS = {
    # question 1 phrasing
    r'capital expenditure': [
        'capital expenditure',
        'revenue expenditure',
        'capital expense',
        'long-term capital expenditure',
    ],
    # question 2 phrasing
    r'(GST input tax credit|input tax credit|ITC)': [
        'GST input tax credit',
        'input tax credit',
        'tax credit for inputs',
    ],
    # numeric values in calculations (simple example)
    r'120,?000': [
        '120,000',
        '100,000',
        '150,000',
        '80,000',
    ],
}

NUMBER_VARIATIONS = [0.7, 0.9, 1.0, 1.1, 1.25]


def choose_replacement(pattern: str, match_text: str) -> str:
    choices = SUBSTITUTIONS.get(pattern, [match_text])
    return random.choice(choices)


def randomize_numbers(text: str) -> str:
    # Find any number groups like 120000 or 120,000 and slightly vary them
    def repl(m):
        s = m.group(0)
        digits = re.sub(r'[^0-9]', '', s)
        try:
            val = int(digits)
        except ValueError:
            return s
        factor = random.choice(NUMBER_VARIATIONS)
        new_val = int(val * factor)
        return f"{new_val:,}"

    return re.sub(r'\d{2,3}(?:,?\d{3})+', repl, text)


def apply_substitutions(text: str) -> str:
    out = text
    for pattern in SUBSTITUTIONS.keys():
        # case-insensitive replacement
        out = re.sub(pattern, lambda m: choose_replacement(pattern, m.group(0)), out, flags=re.IGNORECASE)
    # Randomize numbers
    out = randomize_numbers(out)
    return out


def generate_from_template(template_path: Path, out_dir: Path, count: int = 5, prefix: str = 'student_var'):
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(str(template_path))
    base_texts = [p.text for p in doc.paragraphs]

    for i in range(count):
        new_doc = Document()
        for p in base_texts:
            new_p = apply_substitutions(p)
            new_doc.add_paragraph(new_p)

        out_name = f"{prefix}_{i+1:03}.docx"
        out_path = out_dir / out_name
        new_doc.save(str(out_path))
        print('Wrote', out_path)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--count', type=int, default=10)
    parser.add_argument('--outdir', type=str, default=str(OUT_DIR))
    parser.add_argument('--templates', nargs='+', default=[str(SAMPLES_DIR / 'student_quiz_001.docx')])
    args = parser.parse_args()

    out_path = Path(args.outdir)
    for t in args.templates:
        tpl = Path(t)
        if not tpl.exists():
            print('Template not found:', tpl)
            continue
        generate_from_template(tpl, out_path, count=args.count, prefix=tpl.stem)

    print('Done')
